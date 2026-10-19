"""Render the cited KOA tracker rows as a video-safe graphic from the workbook."""

import argparse
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

NAVY = "#0a1628"
AMBER = "#e8a33d"
PAPER = "#f5efe6"
CYAN = "#c9d6df"


def font(name: str, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(Path("C:/Windows/Fonts") / name), size)


def fit(draw: ImageDraw.ImageDraw, value: str, face: ImageFont.FreeTypeFont, width: int) -> str:
    while draw.textlength(value, font=face) > width:
        value = value[:-2].rstrip("…") + "…"
    return value


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    sheet = load_workbook(args.workbook, read_only=True, data_only=True)["Project Plan"]
    rows = {number: [sheet.cell(number, column).value for column in range(2, 7)] for number in range(27, 32)}
    if rows[29][1] != "Store validated training pairs and consent records" or rows[29][4] != "P0":
        raise ValueError("Project Plan row 29 differs from the requested receipt")

    image = Image.new("RGB", (1920, 1080), NAVY)
    draw = ImageDraw.Draw(image)
    serif = font("georgia.ttf", 66)
    mono = font("consola.ttf", 25)
    mono_bold = font("consolab.ttf", 28)
    body = font("segoeui.ttf", 32)
    body_bold = font("segoeuib.ttf", 34)

    draw.rectangle((0, 0, 1920, 12), fill=AMBER)
    draw.text((114, 80), "KOA WEBSITE v4", font=mono_bold, fill=AMBER)
    draw.text((114, 133), "Project Tracker", font=serif, fill=PAPER)
    draw.text((115, 235), "SOURCE  /  PROJECT PLAN  /  ROW 29", font=mono, fill=CYAN)

    left, right = 112, 1808
    top, row_height = 326, 112
    draw.rounded_rectangle((left, top, right, top + 6 * row_height + 4), radius=18, fill="#10233a", outline="#527187", width=2)
    columns = [(140, "ROW"), (230, "LANE"), (450, "TRACKER TASK"), (1320, "STATUS"), (1590, "PRIORITY")]
    for x, label in columns:
        draw.text((x, top + 29), label, font=mono, fill=CYAN)
    draw.line((left, top + row_height, right, top + row_height), fill="#527187", width=2)

    for index, number in enumerate(range(27, 32)):
        y = top + (index + 1) * row_height
        lane, task, _, status, priority = rows[number]
        if number == 29:
            draw.rounded_rectangle((left + 7, y + 4, right - 7, y + row_height - 4), radius=12, fill="#253044", outline=AMBER, width=5)
        elif index % 2:
            draw.rectangle((left + 8, y + 3, right - 8, y + row_height - 3), fill="#142a42")
        face = body_bold if number == 29 else body
        color = PAPER if number == 29 else "#aebdca"
        draw.text((140, y + 36), str(number), font=mono_bold, fill=AMBER if number == 29 else CYAN)
        draw.text((230, y + 34), str(lane), font=face, fill=color)
        draw.text((450, y + 34), fit(draw, str(task), face, 820), font=face, fill=color)
        draw.text((1320, y + 37), str(status), font=mono, fill=color)
        draw.rounded_rectangle((1590, y + 29, 1662, y + 76), radius=8, fill=AMBER if number == 29 else "#2b4158")
        draw.text((1607, y + 36), str(priority), font=mono_bold, fill=NAVY if number == 29 else CYAN)

    draw.text((115, 1020), "TRACKER STATE: IN PROGRESS  ·  TASK RECORD, NOT PROOF OF COMPLETION", font=mono, fill=CYAN)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    image.save(args.output)
    print(f"{args.output}: {image.width}x{image.height}; Project Plan!C29/F29 verified")


if __name__ == "__main__":
    main()
